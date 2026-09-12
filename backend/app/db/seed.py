"""
Seed con datos REALES del Excel.

1) Coloca el archivo en:
   backend/app/db/data/Torta_Trozada_Cadmio_ordenado.xlsx

2) Ejecuta:
   docker compose exec backend python -m app.db.seed
"""
from __future__ import annotations
from pathlib import Path
from datetime import datetime, date
from app.db.session import SessionLocal, engine, Base
from app.core.security import get_password_hash
from app.models.user import User, UserRole
from app.models.catalog import Product, Origin
from app.models.lot import Lot, LotOrigin
from app.models.sample import SampleLot, SampleGrain

DATA_DIR = Path(__file__).resolve().parent / "data"
EXCEL_NAME = "Torta_Trozada_Cadmio_ordenado.xlsx"


def _excel_path() -> Path:
    p = DATA_DIR / EXCEL_NAME
    if p.exists():
        return p
    for alt in [
        Path("/app/data") / EXCEL_NAME,
        Path("/data") / EXCEL_NAME,
    ]:
        if alt.exists():
            return alt
    raise FileNotFoundError(
        f"No se encontró {EXCEL_NAME}. Cópialo a backend/app/db/data/ y reinicia el backend."
    )


def get_or_create_origin(db, name: str) -> Origin:
    if not name:
        name = "SIN ORIGEN"
    name = " ".join(str(name).strip().split())
    o = db.query(Origin).filter(Origin.name.ilike(name)).first()
    if not o:
        o = Origin(name=name, country="Perú")
        db.add(o)
        db.flush()
    return o


def get_or_create_product(db, name: str) -> Product:
    name = str(name).strip()
    p = db.query(Product).filter(Product.name.ilike(name)).first()
    if not p:
        p = Product(name=name)
        db.add(p)
        db.flush()
    return p


def parse_origins(text) -> list[str]:
    if not text:
        return []
    parts = []
    for chunk in str(text).replace("/", ",").split(","):
        n = chunk.strip()
        if n and n.lower() not in ("0", "no pta *", "no pta"):
            parts.append(n)
    return parts


def parse_weight(w):
    if w is None:
        return None
    if isinstance(w, (int, float)):
        return float(w)
    s = str(w).lower().replace("gr", "").replace("g", "").strip()
    try:
        return float(s)
    except Exception:
        return None


def to_date(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        return date.fromisoformat(str(v)[:10])
    except Exception:
        return None


def _add_lot_sample(db, *, product_name, lot_code, cd, obs, pesticides, send_date, origins, weight_g, producer_code, producer_name):
    product = get_or_create_product(db, product_name)
    lot_code = str(lot_code).strip().rstrip("_")
    lot = db.query(Lot).filter(Lot.lot_code == lot_code).first()
    if not lot:
        lot = Lot(product_id=product.id, lot_code=lot_code)
        db.add(lot)
        db.flush()
        for oname in parse_origins(origins):
            o = get_or_create_origin(db, oname)
            if not db.query(LotOrigin).filter(LotOrigin.lot_id == lot.id, LotOrigin.origin_id == o.id).first():
                db.add(LotOrigin(lot_id=lot.id, origin_id=o.id))

    if db.query(SampleLot).filter(SampleLot.lot_id == lot.id).first():
        return 0

    has = isinstance(cd, (int, float)) and not (obs and "SIN MUESTRA" in str(obs).upper())
    db.add(SampleLot(
        lot_id=lot.id,
        cadmium_mg_kg=float(cd) if has else None,
        has_sample=has,
        sample_weight_g=weight_g,
        pesticides=pesticides,
        observation=obs,
        send_date=to_date(send_date),
        producer_code=str(producer_code) if producer_code else None,
        producer_name=str(producer_name) if producer_name else None,
    ))
    return 1


def seed_from_excel(db):
    import openpyxl

    path = _excel_path()
    wb = openpyxl.load_workbook(path, data_only=True)
    n_samples = 0
    n_grain = 0

    # --- Torta de Cacao ---
    if "Torta de Cacao" in wb.sheetnames:
        ws = wb["Torta de Cacao"]
        for row in ws.iter_rows(min_row=4, values_only=True):
            if not row or not row[4]:
                continue
            product, peso, prod_code, prod_name, lote, cd, obs, pest, fecha, origins = (list(row) + [None] * 10)[:10]
            n_samples += _add_lot_sample(
                db,
                product_name=product or "Torta de cacao",
                lot_code=lote,
                cd=cd,
                obs=obs,
                pesticides=pest,
                send_date=fecha,
                origins=origins,
                weight_g=parse_weight(peso) or 1000.0,
                producer_code=prod_code,
                producer_name=prod_name,
            )

    # --- Torta de cacao alcalino ---
    for sheet in wb.sheetnames:
        if sheet.lower().strip() == "torta de cacao alcalino":
            ws = wb[sheet]
            for row in ws.iter_rows(min_row=4, values_only=True):
                if not row or not row[4]:
                    continue
                cells = list(row)
                product, peso, prod_code, prod_name, lote, cd, obs, fecha, origins = (cells + [None] * 9)[:9]
                n_samples += _add_lot_sample(
                    db,
                    product_name=product or "Torta de cacao alcalino",
                    lot_code=lote,
                    cd=cd,
                    obs=obs,
                    pesticides=None,
                    send_date=fecha,
                    origins=origins,
                    weight_g=parse_weight(peso) or 300.0,
                    producer_code=prod_code,
                    producer_name=prod_name,
                )

    # --- Cacao alcalino reducido en grasa ---
    if "Cacao alcalino" in wb.sheetnames:
        ws = wb["Cacao alcalino"]
        for row in ws.iter_rows(min_row=4, values_only=True):
            if not row or not row[4]:
                continue
            product, peso, prod_code, prod_name, lote, cd, obs, fecha, origins = (list(row) + [None] * 9)[:9]
            n_samples += _add_lot_sample(
                db,
                product_name=product or "Cacao alcalino reducido en grasa",
                lot_code=lote,
                cd=cd,
                obs=obs,
                pesticides=None,
                send_date=fecha,
                origins=origins,
                weight_g=parse_weight(peso) or 300.0,
                producer_code=prod_code,
                producer_name=prod_name,
            )

    # --- Cacao en polvo ---
    if "Cacao en polvo" in wb.sheetnames:
        ws = wb["Cacao en polvo"]
        for row in ws.iter_rows(min_row=4, values_only=True):
            if not row or not row[4]:
                continue
            product, peso, prod_code, prod_name, lote, cd, obs, fecha, origins = (list(row) + [None] * 9)[:9]
            n_samples += _add_lot_sample(
                db,
                product_name=product or "Cacao en polvo",
                lot_code=lote,
                cd=cd,
                obs=obs,
                pesticides=None,
                send_date=fecha,
                origins=origins,
                weight_g=parse_weight(peso) or 300.0,
                producer_code=prod_code,
                producer_name=prod_name,
            )

    # --- Torta trozada (plaguicidas / orígenes; sin cadmio si no está en otra hoja) ---
    if "Torta trozada estándar" in wb.sheetnames:
        ws = wb["Torta trozada estándar"]
        for row in ws.iter_rows(min_row=4, values_only=True):
            if not row or not row[0]:
                continue
            lote, pest, origins, tm = (list(row) + [None] * 4)[:4]
            n_samples += _add_lot_sample(
                db,
                product_name="Torta trozada estándar",
                lot_code=lote,
                cd=None,
                obs=None,
                pesticides=pest,
                send_date=None,
                origins=origins,
                weight_g=300.0,
                producer_code="Chincha",
                producer_name="Exportadora Romex S.A",
            )

    # --- Grano de cacao ---
    if "Grano de cacao" in wb.sheetnames:
        ws = wb["Grano de cacao"]
        for row in ws.iter_rows(min_row=4, values_only=True):
            cells = list(row)
            if all(c is None for c in cells[:5]):
                continue
            product, peso, origen, guia, cd, obs, fecha = (cells + [None] * 7)[:7]
            if not guia and not origen:
                continue
            o = get_or_create_origin(db, origen or "SIN ORIGEN")
            guia_code = str(guia).strip() if guia else ""
            if not guia_code:
                continue
            if db.query(SampleGrain).filter(SampleGrain.origin_id == o.id, SampleGrain.guia_code == guia_code).first():
                continue
            has = isinstance(cd, (int, float)) and not (obs and "SIN MUESTRA" in str(obs).upper())
            db.add(SampleGrain(
                origin_id=o.id,
                guia_code=guia_code,
                cadmium_mg_kg=float(cd) if has else None,
                has_sample=has,
                sample_weight_g=parse_weight(peso) or 350.0,
                observation=obs,
                send_date=to_date(fecha),
            ))
            n_grain += 1

    return n_samples, n_grain


def seed():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        admin = db.query(User).filter(
            (User.email == "admin@cadmio.com") | (User.email == "admin@cadmio.local")
        ).first()
        if not admin:
            db.add(User(
                email="admin@cadmio.com",
                full_name="Administrador",
                hashed_password=get_password_hash("admin123"),
                role=UserRole.ADMIN,
            ))
        elif admin.email == "admin@cadmio.local":
            admin.email = "admin@cadmio.com"
        db.commit()

        n_samples, n_grain = seed_from_excel(db)
        db.commit()
        print(f"Seed REAL desde Excel: {n_samples} muestras producto, {n_grain} muestras grano")
        print("Login: admin@cadmio.com / admin123")
    except Exception as e:
        db.rollback()
        print(f"Error seed: {e}")
        raise
    finally:
        db.close()


if __name__ == "__main__":
    seed()
