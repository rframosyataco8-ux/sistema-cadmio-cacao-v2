"""
Importa el Excel de Lima Cadmio a PostgreSQL.
Uso:
  docker compose cp "lima cadmio (1).xlsx" backend:/tmp/lima.xlsx
  docker compose exec backend python -m app.db.import_excel /tmp/lima.xlsx
"""
from __future__ import annotations
import re
import sys
from datetime import datetime, date
from pathlib import Path

from openpyxl import load_workbook

from app.db.session import SessionLocal, engine, Base
from app.models.catalog import Product, Origin
from app.models.lot import Lot, LotOrigin
from app.models.sample import SampleLot, SampleGrain


def parse_num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", ".")
    s = re.sub(r"[^\d.\-]", "", s)
    if not s or s in (".", "-", "-."):
        return None
    try:
        return float(s)
    except Exception:
        return None


def parse_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    try:
        return datetime.fromisoformat(s[:10]).date()
    except Exception:
        return None


def parse_weight(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = re.search(r"([\d.,]+)", str(v))
    return parse_num(m.group(1)) if m else None


def clean_pest(v):
    if v is None:
        return None
    s = str(v).strip()
    if not s or s.lower() in ("0", "no pta *", "no pta", "none", "-"):
        return None
    return s


def is_no_sample(obs):
    return bool(obs and "sin muestra" in str(obs).lower())


def get_or_create_origin(db, name: str) -> Origin:
    if not name:
        name = "SIN ORIGEN"
    name = " ".join(str(name).strip().split())
    o = db.query(Origin).filter(Origin.name.ilike(name)).first()
    if not o:
        o = Origin(name=name, country="Perú")
        db.add(o)
        db.flush()
    return o


def get_or_create_product(db, name: str) -> Product:
    name = str(name).strip()
    p = db.query(Product).filter(Product.name.ilike(name)).first()
    if not p:
        p = Product(name=name)
        db.add(p)
        db.flush()
    return p


def parse_origins(text) -> list[str]:
    if not text:
        return []
    parts = []
    for chunk in str(text).replace("/", ",").split(","):
        n = chunk.strip()
        if n and n.lower() not in ("0", "no pta *", "no pta"):
            parts.append(n)
    return parts


def parse_excel(path: Path):
    wb = load_workbook(path, data_only=True)
    lots, grain = [], []

    ws = wb["Torta trozada estándar"]
    for row in ws.iter_rows(min_row=4, values_only=True):
        lot_code, pest, origins, _tm = row[0], row[1], row[2], row[3]
        if not lot_code:
            continue
        lots.append({
            "product": "Torta trozada estándar",
            "lot_code": str(lot_code).strip(),
            "origins": str(origins or "").strip() or None,
            "date": None,
            "cd": None,
            "has_sample": True,
            "pesticides": clean_pest(pest),
            "obs": None,
            "weight_g": None,
            "producer_code": None,
            "producer_name": None,
        })

    ws = wb["Torta de Cacao"]
    for row in ws.iter_rows(min_row=4, values_only=True):
        product, peso, prod_code, prod_name, lot_code, cd, obs, pest, fecha, origins = row[:10]
        if not lot_code:
            continue
        cd_val = parse_num(cd)
        no_s = is_no_sample(obs)
        if no_s:
            cd_val = None
        lots.append({
            "product": (str(product).strip() if product else "Torta de cacao") or "Torta de cacao",
            "lot_code": str(lot_code).strip(),
            "origins": str(origins or "").strip() or None,
            "date": parse_date(fecha),
            "cd": cd_val,
            "has_sample": False if no_s else (cd_val is not None),
            "pesticides": clean_pest(pest),
            "obs": str(obs).strip() if obs else None,
            "weight_g": parse_weight(peso),
            "producer_code": str(prod_code).strip() if prod_code else None,
            "producer_name": str(prod_name).strip() if prod_name else None,
        })

    for sheet, default_product in (
        ("torta de cacao alcalino", "Torta de cacao alcalino"),
        ("Cacao alcalino", "Cacao alcalino reducido en grasa"),
        ("Cacao en polvo", "Cacao en polvo"),
    ):
        ws = wb[sheet]
        for row in ws.iter_rows(min_row=4, values_only=True):
            cells = (list(row) + [None] * 9)[:9]
            product, peso, prod_code, prod_name, lot_code, cd, obs, fecha, origins = cells
            if not lot_code:
                continue
            cd_val = parse_num(cd)
            no_s = is_no_sample(obs)
            if no_s:
                cd_val = None
            prod = str(product).strip() if product else default_product
            if not prod:
                prod = default_product
            lots.append({
                "product": prod,
                "lot_code": str(lot_code).strip().rstrip("_"),
                "origins": str(origins or "").strip() or None,
                "date": parse_date(fecha),
                "cd": cd_val,
                "has_sample": False if no_s else (cd_val is not None),
                "pesticides": None,
                "obs": str(obs).strip() if obs else None,
                "weight_g": parse_weight(peso),
                "producer_code": str(prod_code).strip() if prod_code else None,
                "producer_name": str(prod_name).strip() if prod_name else None,
            })

    ws = wb["Grano de cacao"]
    for row in ws.iter_rows(min_row=4, values_only=True):
        cells = (list(row) + [None] * 7)[:7]
        product, peso, origin, guia, cd, obs, fecha = cells
        guia_s = str(guia).strip() if guia else ""
        if not guia_s:
            continue
        origin_s = str(origin).strip() if origin else "SIN ORIGEN"
        if not origin_s:
            origin_s = "SIN ORIGEN"
        cd_val = parse_num(cd)
        no_s = is_no_sample(obs)
        if no_s:
            cd_val = None
        grain.append({
            "origin": origin_s,
            "guia_code": guia_s,
            "cd": cd_val,
            "has_sample": False if no_s else (cd_val is not None),
            "weight_g": parse_weight(peso) or 350,
            "obs": str(obs).strip() if obs else None,
            "date": parse_date(fecha),
            "is_organic": bool(product and "organic" in str(product).lower()),
        })

    return lots, grain


def import_data(path: Path):
    Base.metadata.create_all(bind=engine)
    lots, grain = parse_excel(path)
    db = SessionLocal()
    try:
        n_lots = n_samples_new = n_samples_upd = n_grain = 0
        for row in lots:
            product = get_or_create_product(db, row["product"])
            lot = (
                db.query(Lot)
                .filter(Lot.product_id == product.id, Lot.lot_code == row["lot_code"])
                .first()
            )
            if not lot:
                lot = Lot(product_id=product.id, lot_code=row["lot_code"])
                db.add(lot)
                db.flush()
                n_lots += 1
            for oname in parse_origins(row.get("origins")):
                o = get_or_create_origin(db, oname)
                if not db.query(LotOrigin).filter(
                    LotOrigin.lot_id == lot.id, LotOrigin.origin_id == o.id
                ).first():
                    db.add(LotOrigin(lot_id=lot.id, origin_id=o.id))
            sample = db.query(SampleLot).filter(SampleLot.lot_id == lot.id).first()
            if sample:
                if row.get("cd") is not None:
                    sample.cadmium_mg_kg = row["cd"]
                    sample.has_sample = bool(row.get("has_sample", True))
                if row.get("pesticides") is not None:
                    sample.pesticides = row["pesticides"]
                if row.get("obs"):
                    sample.observation = row["obs"]
                if row.get("date") and not sample.send_date:
                    sample.send_date = row["date"]
                if row.get("weight_g") and not sample.sample_weight_g:
                    sample.sample_weight_g = row["weight_g"]
                n_samples_upd += 1
            else:
                db.add(SampleLot(
                    lot_id=lot.id,
                    cadmium_mg_kg=row.get("cd"),
                    has_sample=bool(row.get("has_sample", True)),
                    sample_weight_g=row.get("weight_g"),
                    pesticides=row.get("pesticides"),
                    observation=row.get("obs"),
                    send_date=row.get("date"),
                    producer_code=row.get("producer_code"),
                    producer_name=row.get("producer_name"),
                ))
                n_samples_new += 1

        for row in grain:
            o = get_or_create_origin(db, row.get("origin") or "SIN ORIGEN")
            guia = str(row.get("guia_code") or "").strip()
            if not guia:
                continue
            existing = db.query(SampleGrain).filter(
                SampleGrain.origin_id == o.id, SampleGrain.guia_code == guia
            ).first()
            if existing:
                if row.get("cd") is not None:
                    existing.cadmium_mg_kg = row["cd"]
                    existing.has_sample = bool(row.get("has_sample", True))
                continue
            db.add(SampleGrain(
                origin_id=o.id,
                guia_code=guia,
                cadmium_mg_kg=row.get("cd"),
                has_sample=bool(row.get("has_sample", True)),
                sample_weight_g=row.get("weight_g"),
                observation=row.get("obs"),
                send_date=row.get("date"),
                is_organic=bool(row.get("is_organic", False)),
            ))
            n_grain += 1

        db.commit()
        print(
            f"Import OK: {n_lots} lotes nuevos, {n_samples_new} muestras nuevas, "
            f"{n_samples_upd} actualizadas, {n_grain} grano nuevos | total excel lotes={len(lots)} grain={len(grain)}"
        )
        pending = [r["lot_code"] for r in lots if r["product"] == "Torta de cacao" and r["cd"] is None and not (r.get("obs") and "sin muestra" in str(r["obs"]).lower())]
        print("Pendientes análisis (torta cacao sin Cd):", pending)
    except Exception as e:
        db.rollback()
        print(f"Error import: {e}")
        raise
    finally:
        db.close()


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Uso: python -m app.db.import_excel /ruta/al/archivo.xlsx")
        sys.exit(1)
    p = Path(sys.argv[1])
    if not p.exists():
        print(f"No existe: {p}")
        sys.exit(1)
    import_data(p)
